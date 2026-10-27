"""Build an editable Tajik competition deck from actual OBHUSH interface captures."""
from pathlib import Path
import json, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from reportlab.graphics.barcode.qr import QrCodeWidget
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/presentation'
ASSETS = OUT / 'assets'
CONTENT = json.loads((OUT/'content-tj.json').read_text())
P = Presentation()
P.slide_width, P.slide_height = Inches(13.333333), Inches(7.5)
C = dict(forest='183C31', ink='203C32', muted='64776B', cream='F7F5EF', lime='D7EAA5', white='FFFFFF', pale='E9EEE2', teal='237F86', line='D9E0D3', sand='EEE3CB')
FONT='Noto Sans'
SERIF='Noto Serif'


def color(value):
    return RGBColor.from_string(C.get(value, value))


def rect(s,x,y,w,h,fill='white',radius=False,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=color(fill)
    if radius:
        sh.adjustments[0]=.12
    if line:
        sh.line.color.rgb=color(line);sh.line.width=Pt(.8)
    else:sh.line.fill.background()
    return sh


def circle(s,x,y,d,fill):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid();sh.fill.fore_color.rgb=color(fill);sh.line.fill.background();return sh


def text(s,x,y,w,h,value,size=18,fill='ink',bold=False,font=FONT,align=None):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(fill)
        p.space_before=Pt(0);p.space_after=Pt(5);p.line_spacing=1.12
        if align is not None:p.alignment=align
    return sh


def line(s,x1,y1,x2,y2,fill='line',width=1):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=color(fill);sh.line.width=Pt(width);return sh


def photo(s,path,x,y,w,h,cover=False):
    with Image.open(path) as im:iw,ih=im.size
    if cover:
        pic=s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
        ar,box=iw/ih,w/h
        if ar>box:
            crop=(1-box/ar)/2;pic.crop_left=pic.crop_right=crop
        else:
            crop=(1-ar/box)/2;pic.crop_top=pic.crop_bottom=crop
        return pic
    scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale
    return s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))


def wave(s,x,y,fill='lime',scale=1):
    for k in range(3):
        last=None
        for i in range(25):
            xx=x+i*.023*scale;yy=y+k*.10*scale+math.sin(i*math.pi/12)*.045*scale
            if last:line(s,*last,xx,yy,fill,1.5)
            last=xx,yy


def footer(s,n,dark=False):
    col='lime' if dark else 'muted'
    line(s,.65,7.04,12.68,7.04,'muted' if dark else 'line',.65)
    text(s,.65,7.15,8,.2,'ОБҲУШ  /  Идеяи ман — калиди фардо  /  2026',9,col)
    text(s,11.7,7.13,1,.25,f'{n:02d} / 10',10,col,align=PP_ALIGN.RIGHT)


def slide(n,eyebrow,title,dark=False):
    s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color('forest' if dark else 'cream')
    text(s,.65,.42,11,.25,eyebrow.upper(),10,'lime' if dark else 'muted',True)
    text(s,.65,.98,12,1.03,title,31,'white' if dark else 'ink',font=SERIF)
    footer(s,n,dark)
    s.notes_slide.notes_text_frame.text=CONTENT[n-1]['notes']
    return s


def label(s,x,y,w,value,dark=False):
    rect(s,x,y,w,.37,'lime' if dark else 'pale',True)
    text(s,x+.12,y+.075,w-.24,.2,value,9,'forest',True)


def qr(s,x,y,size):
    widget=QrCodeWidget('https://obihush.site',barLevel='M')
    widget.qr.make();matrix=widget.qr.modules;n=len(matrix);unit=size/(n+8)
    rect(s,x,y,size,size,'white')
    for r,row in enumerate(matrix):
        for c,black in enumerate(row):
            if black:rect(s,x+(c+4)*unit,y+(r+4)*unit,unit+.001,unit+.001,'forest')


# 01: branded cover, real supplied gardener still.
s=slide(1,'Номинатсияи 5 · Зеҳни сунъӣ ва таҳлили рақамӣ','',True)
photo(s,ROOT/'public/media/garden-desktop.jpg',7.65,0,5.69,7.03,True)
wave(s,.72,1.28,scale=1.65)
text(s,1.95,1.15,5,.7,'ОБҲУШ',38,'white',True)
text(s,.65,2.35,6.8,1.75,'Об. Обу ҳаво.\nҚарори дуруст.',36,'lime',font=SERIF)
text(s,.68,4.38,6.2,1.0,'Ёвари рақамӣ барои обёрӣ,\nкори берунӣ ва сарфаи об.',20,'white')
label(s,.68,5.94,3.25,'БАРОИ МАРДУМИ ТОҶИКИСТОН',True)
text(s,8.02,6.51,5,.28,'Тасвир аз видеои пешниҳодшудаи лоиҳа',9,'white')

# 02: people before technology.
s=slide(2,'01 · Мушкилот','Маълумот ҳаст. Қарор душвор аст.')
for i,(name,question,desc,fill) in enumerate([
 ('Боғбон','Кай ва чӣ қадар\nоб диҳам?','Боришро чӣ гуна ба обёрӣ\nпайваст кунам?','pale'),
 ('Корманд','Кай дар берун\nкор кунам?','Соатҳои муносибро аз куҷо\nпайдо кунам?','sand'),
 ('Оила','Чӣ қадар об\nсарфа кунам?','Кадом одати ҳаррӯза\nтағйир ёбад?','white')]):
 x=.65+i*4.08;rect(s,x,2.36,3.88,3.61,fill,True)
 text(s,x+.27,2.63,3.3,.35,f'0{i+1}  /  {name}',13,'muted',True)
 text(s,x+.27,3.28,3.32,1.22,question,24,font=SERIF)
 text(s,x+.27,4.96,3.32,.8,desc,16,'muted')
text(s,.7,6.33,12,.45,'Ба инсон амали фаҳмо лозим аст.',22,'ink',True)

# 03: actual product screenshot and three short steps.
s=slide(3,'02 · Роҳи ҳал','Шаҳри худ. Вазифаи худ. Нақшаи худ.')
photo(s,ASSETS/'weather-tj.png',.65,2.18,8.25,4.4)
rect(s,9.22,2.25,3.46,4.27,'forest',True)
for i,(head,desc) in enumerate([('Ҷойи худ','12 шаҳр ё ҷойгиршавӣ'),('Вазифаи худ','Боғ, кор ё оила'),('Нақшаи худ','Ҳисоб, шарҳ ва нигоҳдорӣ')]):
 y=2.65+i*1.13
 text(s,9.5,y,.6,.35,f'0{i+1}',16,'lime',True)
 text(s,10.1,y,2.3,.4,head,15,'white',True)
 text(s,10.1,y+.45,2.25,.5,desc,12,'lime')
text(s,.72,6.58,8,.2,'Тасвири интерфейси воқеӣ · шароити ҳозира ва пешгӯии рӯз ҷудоанд',10,'muted')

# 04: gardening.
s=slide(4,'03 · Боғ ва обёрӣ','Барои боғбон: обёрӣ бо ҳисоб.')
text(s,.7,2.23,4.05,.65,'Параметрҳоро ворид кун.',20,'ink',True)
for i,(a,b) in enumerate([('Масоҳат','Қитъаи худ, бо м²'),('Зироат ва усул','Коэффисиент ва самаранокӣ'),('Бориш ва бухоршавӣ','Аз пешгӯии ҷои интихобшуда')]):
 y=3.08+i*.92;circle(s,.73,y+.05,.12,'teal');text(s,1.0,y,3.67,.3,a,17,bold=True);text(s,1.0,y+.38,3.67,.35,b,13,'muted')
photo(s,ASSETS/'garden-tj.png',4.82,2.1,7.86,4.54)
rect(s,.72,6.2,3.75,.47,'pale',True);text(s,.85,6.3,3.5,.25,'Натиҷа: арзёбии об бо литр',12,'ink',True)
text(s,4.95,6.62,7.6,.21,'Арзёбии ҳисобӣ; намии воқеии хок ҳоло чен намешавад.',10,'muted')

# 05: outdoor work.
s=slide(5,'04 · Кор дар берун','Барои корманд: вақти муносибро интихоб кун.')
photo(s,ASSETS/'work-tj.png',.65,2.06,8.07,4.62)
text(s,9.1,2.49,3.5,.7,'Аз рӯйи шароит,\nна танҳо соат.',22,font=SERIF)
for i,v in enumerate(['Рӯшноии рӯз','Ҳарорат ва шамол','Бориш ва осмон']):
 text(s,9.13,3.68+i*.57,3.4,.38,'— '+v,16,'muted')
label(s,9.11,5.63,3.46,'СОДИРОТ БА ТАҚВИМ · ICS')
text(s,.73,6.62,11.8,.2,'Интихоби соатҳои прогнозӣ; дастурҳои бехатарии кори мушаххас риоя мешаванд.',10,'muted')

# 06: mathematically correct illustrative family scenario.
s=slide(6,'05 · Хона ва оила','Одати хурд. Сарфаи намоён.')
rect(s,.65,2.25,6.02,4.35,'white',True)
text(s,.96,2.59,5.4,.38,'МИСОЛИ ҲИСОБӢ',12,'muted',True)
for y,title,val,width,fill in [(3.4,'Ҳозир: 8 дақиқа','288 л / рӯз',4.85,'muted'),(4.58,'Пас аз тағйир: 5 дақиқа','180 л / рӯз',3.03,'teal')]:
 text(s,.99,y,3.3,.35,title,17,bold=True);text(s,4.3,y,1.94,.35,val,14,'muted',align=PP_ALIGN.RIGHT);rect(s,1.0,y+.52,width,.24,fill,True)
text(s,.99,5.81,5.4,.42,'4 нафар × 9 л/дақ · танҳо истифодаи душ',12,'muted')
text(s,7.18,2.59,5.2,.4,'ИМКОНИ САРФА ДАР 30 РӮЗ',12,'muted',True)
text(s,7.1,3.27,5.5,1.3,'3 240',68,'forest',True)
text(s,7.22,4.76,4.9,.75,'литр об',31,'teal',font=SERIF)
text(s,7.22,5.88,5.0,.6,'108 л / рӯз × 30 рӯз',18,'muted')
text(s,.72,6.73,12,.2,'Модели одатҳо бо параметрҳои интихобшуда; сарфаи воқеӣ бояд бо ҳисобкунак санҷида шавад.',10,'muted')

# 07: real retained river and safety scope.
s=slide(7,'06 · Об, дарё ва омилҳои хатар','Аз қитъаи хона то ҳавзаи Вахш.')
photo(s,ASSETS/'rivers-tj.png',.65,2.15,7.4,4.7)
for i,(h,b) in enumerate([('Харита ва релеф','Норак, Рашт ва болооб'),('Ҷараён ва сенария','Пешгӯӣ, НБО ва обёрӣ'),('Бориш ва таваҷҷуҳ','Шарҳи сабаби аломат дар харита')]):
 y=2.48+i*1.08;text(s,8.5,y,4.0,.4,h,19,bold=True);text(s,8.5,y+.47,4.0,.42,b,13,'muted')
rect(s,8.45,5.91,4.2,.65,'sand',True);text(s,8.63,6.04,3.84,.42,'Барои пешгӯии сел: маълумот ва\nсанҷиши маҳаллӣ ҳанӯз заруранд.',11,'ink')

# 08: honest architecture with clear role separation.
s=slide(8,'07 · Зеҳни сунъӣ ва таҳлили рақамӣ','Маълумот → ҳисоб → шарҳи фаҳмо',True)
for x,title,sub in [(.65,'МАНБАЪҲО','Обу ҳаво, дарё ва релеф'),(4.73,'МУҲАРРИК','Модел ва қоидаҳои санҷишпазир'),(8.81,'НАТИҶА','Нақша барои корбар')]:
 rect(s,x,2.27,3.87,3.96,'cream',True);text(s,x+.25,2.55,3.38,.33,title,12,'muted',True);text(s,x+.25,3.05,3.36,.9,sub,22,'ink',True)
text(s,.9,4.42,3.3,1.3,'Open-Meteo · обу ҳаво\nGloFAS · ҷараёни моделӣ\nAWS / Mapzen · релеф',15,'muted')
text(s,4.98,4.42,3.32,1.4,'Ridge · пешгӯии ҷараён\nФормула · литр ва сарфа\nҚоида · соатҳои кор',15,'muted')
text(s,9.06,4.42,3.32,1.4,'Ҷой + рӯз + вазифа\nЁвари забонии ИИ\nШарҳ ба тоҷикӣ ва русӣ',15,'muted')
text(s,.73,6.55,12,.25,'Санҷиши Ridge: 14 рӯзи ҷудошуда · маълумоти GloFAS моделӣ аст, на ченкунии гидропост.',10,'lime')

# 09: completed evidence; QR links, without claiming deployment of newest interface.
s=slide(9,'08 · Ҳолати лоиҳа','Нусхаи корӣ, ки метавон нишон дод.')
for i,(num,lab,sub) in enumerate([('3','вазифаи ҳаррӯза','Боғ · кор · оила'),('2','забони интерфейс','Тоҷикӣ · русӣ'),('12','шаҳр дар каталог','Инчунин интихоби ҷойгиршавӣ')]):
 x=.7+i*2.61;text(s,x,2.48,2.3,.92,num,51,'forest',True);text(s,x,3.61,2.4,.6,lab,17,bold=True);text(s,x,4.35,2.35,.75,sub,12,'muted')
rect(s,.7,5.48,7.35,.94,'pale',True);text(s,.94,5.69,6.9,.55,'32 санҷиши ҳисоб / API\n16 сенарияи браузерӣ — гузаштанд',14,'ink',True)
rect(s,8.78,2.27,3.88,4.36,'white',True);qr(s,9.18,2.49,3.07)
link=text(s,9.08,5.7,3.3,.5,'obihush.site',22,'teal',True,align=PP_ALIGN.CENTER)
link.text_frame.paragraphs[0].runs[0].hyperlink.address='https://obihush.site'
text(s,9.13,6.23,3.2,.22,'Кушодани сомонаи лоиҳа',10,'muted',align=PP_ALIGN.CENTER)

# 10: strong close, clear next steps.
s=slide(10,'09 · Рушд ва таъсир','',True)
wave(s,.73,1.17,scale=1.45);text(s,1.82,1.03,5,.6,'ОБҲУШ',30,'white',True)
text(s,.67,2.17,8.6,1.68,'Ҳар қатраи об —\nяк қарори беҳтар.',37,'lime',font=SERIF)
text(s,.72,4.25,7.85,.8,'Қадами навбатӣ: озмоиш бо оилаҳо ва боғбонон,\nченкунии сарф ва мутобиқсозии ҳисобҳо.',19,'white')
for i,(a,b) in enumerate([('01','Озмоиши маҳаллӣ'),('02','Ченкунии натиҷа'),('03','Пайвасти датчикҳо')]):
 x=.75+i*2.75;text(s,x,5.83,.45,.3,a,12,'lime',True);text(s,x+.56,5.81,2.12,.65,b,13,'white')
rect(s,9.37,1.9,3.29,4.32,'cream',True);qr(s,9.57,2.14,2.89)
text(s,9.55,5.19,2.93,.5,'obihush.site',21,'forest',True,align=PP_ALIGN.CENTER)
text(s,9.53,5.77,2.97,.25,'Ташаккур барои таваҷҷуҳ!',12,'ink',align=PP_ALIGN.CENTER)
P.core_properties.title='ОБҲУШ — Об. Обу ҳаво. Қарори дуруст.'
P.core_properties.subject='Идеяи ман — калиди фардо · Номинатсияи 5'
P.core_properties.author='ОБҲУШ'
P.core_properties.keywords='ОБҲУШ, Тоҷикистон, об, зеҳни сунъӣ, Canva'
P.core_properties.comments='Editable shapes and text. Real interface captures. Speaker notes in Tajik.'
P.save(OUT/'OBHUSH-TJ.pptx')
notes=['# Матни баромад · ОБҲУШ','', 'Барои 10 слайд, тақрибан 5–6 дақиқа бо намоиши кӯтоҳ.','']
for i,item in enumerate(CONTENT,1):notes += [f'## {i}. {item["title"]}', '',item['notes'],'']
notes += ['## Пеш аз баромад','', 'Версияи навро дар VPS навсозӣ ва санҷед. QR-код сомонаи obihush.site-ро мекушояд; тасвирҳои слайдҳо аз версияи нави омодашуда гирифта шудаанд. Агар навсозӣ ҳанӯз анҷом нашуда бошад, намоиши зиндаи интерфейси навро ваъда надиҳед.','', 'Мисоли 3240 литр сенарияи ҳисобӣ аст. Сарфаи воқеии истифодабарандагон ҳанӯз чен нашудааст. Огоҳсозии расмии сел, датчикҳои маҳаллӣ ва гидропостҳои зинда дар ин версия пайваст нестанд.']
(OUT/'SPEAKER-NOTES-TJ.md').write_text('\n'.join(notes)+'\n')
print(OUT/'OBHUSH-TJ.pptx')
